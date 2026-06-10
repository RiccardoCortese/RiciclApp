import json, os, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

d = json.load(open('test-results.json', encoding='utf-8'))

FONT = 'Arial'
HEADER_FILL = PatternFill('solid', fgColor='2E7D32')   # verde RiciclApp
HEADER_FONT = Font(name=FONT, bold=True, color='FFFFFF', size=11)
TITLE_FONT = Font(name=FONT, bold=True, size=16, color='1B5E20')
SUB_FONT = Font(name=FONT, italic=True, size=10, color='555555')
CELL_FONT = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
PASS_FILL = PatternFill('solid', fgColor='C8E6C9')
FAIL_FILL = PatternFill('solid', fgColor='FFCDD2')
THIN = Side(style='thin', color='BBBBBB')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center')
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)


def suite_name(path):
    return os.path.basename(path.replace('\\', '/'))


def clean(s):
    return s.replace('—', '-').replace('–', '-')


rows = []
per_suite = {}
for s in d['testResults']:
    sn = suite_name(s['name'])
    for a in s['assertionResults']:
        group = clean(a['ancestorTitles'][0]) if a['ancestorTitles'] else ''
        rows.append({
            'suite': sn,
            'group': group,
            'test': clean(a['title']),
            'status': a['status'],
            'duration': a.get('duration') or 0,
        })
        st = per_suite.setdefault(sn, {'passed': 0, 'failed': 0, 'dur': 0})
        st['passed'] += 1 if a['status'] == 'passed' else 0
        st['failed'] += 1 if a['status'] != 'passed' else 0
        st['dur'] += a.get('duration') or 0

run_dt = datetime.datetime.fromtimestamp(d['startTime'] / 1000).strftime('%d/%m/%Y %H:%M')
total = d['numTotalTests']
passed = d['numPassedTests']
failed = d['numFailedTests']
pass_rate = (passed / total * 100) if total else 0

wb = Workbook()

# ── Foglio 1: Riepilogo ──────────────────────────────────────────────────────
ws = wb.active
ws.title = 'Riepilogo'
ws.sheet_view.showGridLines = False

ws['A1'] = 'RiciclApp - Risultati Test Frontend'
ws['A1'].font = TITLE_FONT
ws['A2'] = f'Component testing (Jest + React Native Testing Library) - Eseguito il {run_dt}'
ws['A2'].font = SUB_FONT

summary = [
    ('Suite di test', d['numTotalTestSuites']),
    ('Test totali', total),
    ('Superati (PASS)', passed),
    ('Falliti (FAIL)', failed),
    ('Tasso di successo', f'{pass_rate:.1f}%'),
]
r = 4
for label, val in summary:
    ws.cell(r, 1, label).font = BOLD
    c = ws.cell(r, 2, val)
    c.font = CELL_FONT
    if label == 'Superati (PASS)':
        c.fill = PASS_FILL
    if label == 'Falliti (FAIL)' and val:
        c.fill = FAIL_FILL
    r += 1

r += 1
ws.cell(r, 1, 'Dettaglio per suite (screen)').font = Font(name=FONT, bold=True, size=12, color='1B5E20')
r += 1
heads = ['Suite', 'Superati', 'Falliti', 'Totale', 'Durata (ms)']
for j, h in enumerate(heads, 1):
    c = ws.cell(r, j, h)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = CENTER
    c.border = BORDER
r += 1
for sn in sorted(per_suite):
    st = per_suite[sn]
    vals = [sn, st['passed'], st['failed'], st['passed'] + st['failed'], st['dur']]
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v)
        c.font = CELL_FONT
        c.border = BORDER
        c.alignment = LEFT if j == 1 else CENTER
    ws.cell(r, 2).fill = PASS_FILL
    if st['failed']:
        ws.cell(r, 3).fill = FAIL_FILL
    r += 1
tot_dur = sum(st['dur'] for st in per_suite.values())
for j, v in enumerate(['TOTALE', passed, failed, total, tot_dur], 1):
    c = ws.cell(r, j, v)
    c.font = BOLD
    c.border = BORDER
    c.alignment = LEFT if j == 1 else CENTER

for col, w in zip('ABCDE', [30, 12, 10, 10, 13]):
    ws.column_dimensions[col].width = w

# ── Foglio 2: Dettaglio Test ─────────────────────────────────────────────────
ws2 = wb.create_sheet('Dettaglio Test')
ws2.sheet_view.showGridLines = False
heads2 = ['#', 'Suite (screen)', 'Gruppo', 'Test case', 'Esito', 'Durata (ms)']
for j, h in enumerate(heads2, 1):
    c = ws2.cell(1, j, h)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = CENTER
    c.border = BORDER
for i, row in enumerate(rows, 1):
    esito = 'PASS' if row['status'] == 'passed' else 'FAIL'
    vals = [i, row['suite'], row['group'], row['test'], esito, row['duration']]
    for j, v in enumerate(vals, 1):
        c = ws2.cell(i + 1, j, v)
        c.font = CELL_FONT
        c.border = BORDER
        c.alignment = CENTER if j in (1, 5, 6) else LEFT
    ws2.cell(i + 1, 5).fill = PASS_FILL if esito == 'PASS' else FAIL_FILL
    ws2.cell(i + 1, 5).font = BOLD
ws2.freeze_panes = 'A2'
ws2.auto_filter.ref = f'A1:F{len(rows) + 1}'
for col, w in zip('ABCDEF', [5, 20, 42, 58, 9, 12]):
    ws2.column_dimensions[col].width = w

out = 'Results_Frontend.xlsx'
wb.save(out)
print('Saved', out, '-', total, 'tests,', passed, 'passed,', failed, 'failed')
